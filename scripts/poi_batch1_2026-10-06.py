#!/usr/bin/env python3
"""POI batch 1 (2026-10-06): Farm Stands + Adventures (+2 Arts & Entertainment) for the Navigators Guide.
Coordinates/addresses/phones from Google Places lookups made today (place_id recorded in Notes).
Hours: official site where fetched (markets), else Google Places listing as shown today. Descriptions are
assistant drafts from official pages / place type — never from reviews. Existence-guard on normalised name."""
import sys, re
from openpyxl import load_workbook
WB = sys.argv[1]
D = "2026-10-06"
H = ["Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"]

def hrs(*cells):  # 7 cells
    assert len(cells) == 7
    return dict(zip(H, cells))

def row(sheet, name, tags, addr, lat, lng, desc, hours, site="", phone="", notes="", anchor="premise",
        s_start=None, s_end=None, display="Yes", source="Google Places lookup " + D):
    return dict(sheet=sheet, Name=name, Tags=tags, Address=addr, Latitude=lat, Longitude=lng, Description=desc,
                **hours, Website=site, Phone=phone, Notes=f"{D}: added by atlas assistant; {notes}".rstrip("; "),
                Display=display, **{"Anchor Type": anchor, "Season Start": s_start, "Season End": s_end, "Source": source})

ALL = hrs("Open 24 hours","Open 24 hours","Open 24 hours","Open 24 hours","Open 24 hours","Open 24 hours","Open 24 hours")
NONE = hrs("", "", "", "", "", "", "")

ROWS = [
 # ---------------- Farm Stands
 row("Farm Stands", "Lafayette Farmers Market", "Farm Stand; Farmers Market; Saturday; Downtown",
     "N 5th St between Columbia & Ferry, and Main St between 4th & 6th, Lafayette, IN 47901", 40.4186462, -86.891859,
     "Downtown Lafayette's Saturday market, in its 187th season in 2026 — about 120 vendors along Fifth and Main, rain or shine. Run by Greater Lafayette Commerce.",
     hrs("", "", "", "", "", "8am-12:30pm", ""), "https://lafayettefarmersmarket.com/", "+1 765-742-4044",
     "hours + 2026 opening (May 2) from Greater Lafayette Commerce news release; Google's weekday hours are junk (6am–4am) and ignored; season end 'October' — 10-31 is a placeholder, confirm the last Saturday; coords Google place_id ChIJZTA-NkUdE4gRFNhssWjRebg",
     s_start="05-02", s_end="10-31", source="lafayettefarmersmarket.com / GLC release 2026; Google Places " + D),
 row("Farm Stands", "West Lafayette Farmers Market", "Farm Stand; Farmers Market; Wednesday",
     "Cumberland Park, 3065 N Salisbury St, West Lafayette, IN 47906", 40.4628562, -86.9160842,
     "The city's Wednesday-evening market at Cumberland Park, first Wednesday in May through the last Wednesday in October. Run by West Lafayette Parks & Recreation.",
     hrs("", "", "3:30pm-7:30pm", "", "", "", ""), "https://www.westlafayette.in.gov/government/parks-and-recreation/places/farmers-market", "+1 765-775-5110",
     "hours/season from westlafayette.in.gov (Google shows 3:30–7pm; city page says 7:30 — city wins); season bounds are 'first Wed of May / last Wed of Oct' — set 05-01/10-31; place_id ChIJZ-y4v3H9EogR464b4JzcRNM",
     s_start="05-01", s_end="10-31", source="westlafayette.in.gov; Google Places " + D),
 row("Farm Stands", "Purdue Farmers Market", "Farm Stand; Farmers Market; Thursday; Campus",
     "Memorial Mall, Purdue University, West Lafayette, IN 47907", 40.4257807, -86.9143117,
     "Thursday lunchtime market on Purdue's Memorial Mall, May 7 – October 29 in 2026 (no market May 21 or July 2). Produce, plants, baked goods and food vendors.",
     hrs("", "", "", "11am-3pm", "", "", ""), "https://www.purdue.edu/operations/sustainable-ops/initiatives/farmers-market", "+1 765-494-2481",
     "season/hours/exceptions from purdue.edu market page; place_id ChIJ__OHQLLiEogRdT81HjpfzFE",
     s_start="05-07", s_end="10-29", source="purdue.edu; Google Places " + D),
 row("Farm Stands", "Flint Creek Orchard", "Farm Stand; Orchard; U-Pick; Apples; Berries",
     "5405 W 500 S, West Point, IN 47992", 40.3446421, -87.0076572,
     "Family u-pick orchard southwest of Lafayette — strawberries, blackberries, peaches, apples, Asian pears and pumpkins through the season.",
     hrs("9am-11am, 5pm-7pm","9am-11am, 5pm-7pm","9am-11am, 5pm-7pm","9am-11am, 5pm-7pm","9am-11am, 5pm-7pm","9am-12pm","Closed"), "", "+1 765-413-5527",
     "hours as listed on Google today (split hours — confirm with the farm; seasonal); no website found; place_id ChIJ29AYFIXnEogR33R3AMhSxRE"),
 row("Farm Stands", "Evans Acres Orchard", "Farm Stand; Orchard; Cider; Cash Only",
     "1050 W 800 S, Lafayette, IN 47909", 40.3005662, -86.931811,
     "Small orchard south of Lafayette known for apple cider; cash sales.",
     hrs("Closed","1pm-6pm","Closed","1pm-6pm","Closed","9am-1pm","Closed"), "", "+1 765-418-9467",
     "hours as listed on Google today; thin listing — confirm it is open to the public this season; place_id ChIJt9bKBQDhEogRTC0on8JVgPQ"),
 row("Farm Stands", "Annie's Orchard", "Farm Stand; Orchard; U-Pick; Strawberries; Peaches; Blueberries; Cash Only",
     "9624 W 75 S, West Lafayette, IN 47906", 40.406506, -87.0861479,
     "U-pick strawberries, blueberries and peaches west of West Lafayette; cash preferred.",
     NONE, "", "+1 765-430-0405",
     "no hours posted on Google — seasonal u-pick, get hours from the farm before showing as open; place_id ChIJpR6ThLzvEogRe-o4WuRulYg"),
 row("Farm Stands", "Prairie View Farms", "Farm Stand; Farm; Eggs; Sweet Corn",
     "7765 W State Rd 18, Brookston, IN 47923", 40.6068832, -87.0164942,
     "Farm west of Brookston selling eggs, sweet corn and baked goods in season.",
     hrs("8am-4pm","8am-4pm","8am-4pm","8am-4pm","8am-4pm","Closed","Closed"), "", "+1 765-563-6111",
     "UNCONFIRMED whether this is a retail farm stand or a working farm/equipment business — hidden until Laurie checks; place_id ChIJFTQoHC5WEogRm3LQShnx_zo", display="No"),

 # ---------------- Adventures
 row("Adventures", "Prophetstown State Park", "Adventures; State Park; Prairie; Hiking; Biking; Camping; Aquatic Center; The Farm",
     "5545 Swisher Rd, West Lafayette, IN 47906", 40.4973299, -86.842217,
     "State park where the Tippecanoe meets the Wabash: restored tallgrass prairie, fens and floodplain forest, paved and prairie trails, a campground, an aquatic center, and The Farm at Prophetstown (1920s working farm, separate admission). Office: 4112 SR 225 E. Gate fee.",
     hrs("7am-11pm","7am-11pm","7am-11pm","7am-11pm","7am-11pm","7am-11pm","7am-11pm"), "https://www.in.gov/dnr/state-parks/parks-lakes/prophetstown-state-park/", "+1 765-567-4919",
     "the guide's centre point; hours as listed on Google (DNR gate hours — confirm seasonal); place_id ChIJWZmo7McDE4gRLIrIS6G0ZE8", source="in.gov/dnr; Google Places " + D),
 row("Adventures", "Wolf Park", "Adventures; Wildlife; Wolves; Bison; Foxes; Nonprofit; Guided Tours",
     "4004 E 800 N, Battle Ground, IN 47920", 40.533986, -86.830575,
     "Nonprofit wolf research and education facility north of Battle Ground — guided tours of the wolf, fox and bison enclosures; not a zoo. Admission.",
     hrs("1pm-6pm","1pm-6pm","Closed","1pm-6pm","1pm-6pm","9:30am-5:30pm","1pm-6pm"), "https://wolfpark.org/", "+1 765-567-2265",
     "hours as listed on Google today — seasonal, confirm on wolfpark.org; place_id ChIJl3fXrD8BE4gRIRjf7HE2F1M"),
 row("Adventures", "Fort Ouiatenon Historical Park", "Adventures; Historic Site; County Park; Wabash River; Feast of the Hunters' Moon",
     "3129 S River Rd, West Lafayette, IN 47906", 40.4063479, -86.9635938,
     "County park on the Wabash below West Lafayette with a 1930 replica blockhouse near the site of the 1717 French trading post and Wea villages; home of the Feast of the Hunters' Moon each October. The original fort site is on private land nearby.",
     ALL, "https://www.tippecanoehistory.org/", "+1 765-476-8411",
     "Google lists 24 hours — county park, dawn–dusk more likely; confirm with Tippecanoe County Parks; place_id ChIJTZWRfWvjEogROvOKjxvotws"),
 row("Adventures", "Tippecanoe Battlefield Park", "Adventures; Historic Site; County Park; Battlefield; Monument; Trails; Wabash Heritage Trail",
     "200 Battleground Ave, Battle Ground, IN 47920", 40.5068922, -86.8438448,
     "Site of the 7 November 1811 battle between Harrison's force and Prophetstown's defenders; the 1908 monument, grounds along Burnett Creek, and the northern trailhead of the Wabash Heritage Trail. The museum is a separate listing.",
     NONE, "https://www.tippecanoehistory.org/", "+1 765-567-2147",
     "park grounds — hours not posted separately from the museum; coords are the museum's place_id ChIJGSq7vaADE4gR8e4WAEXHarw (same address)"),
 row("Adventures", "Wabash Heritage Trail", "Adventures; Trail; Hiking; Wabash River; Multi-Use",
     "Trailheads: Fort Ouiatenon, Tapawingo Park, Davis Ferry Park, Tippecanoe Battlefield", 40.4270319, -86.9006445,
     "About 13 miles from Fort Ouiatenon to Tippecanoe Battlefield, hugging the Wabash to Davis Ferry and then Burnett Creek; paved near the cities, natural surface beyond. Pin is the Tapawingo/West Lafayette riverfront section.",
     hrs("6:30am-9pm","7am-9:30pm","6:30am-9pm","7am-9:30pm","6:30am-9pm","6:30am-9pm","7am-9:30pm"), "", "+1 765-567-6218",
     "a linear trail; anchor approximate; hours as Google lists them (park hours); place_id ChIJy9cHOCEdE4gRgWgZu4hI2Vs", anchor="approximate"),
 row("Adventures", "Wabash Heritage Trail — Battle Ground access", "Adventures; Trailhead; Hiking; Burnett Creek",
     "Prophetstown Rd near the Battlefield, West Lafayette, IN 47906", 40.504262, -86.8475821,
     "Northern access to the Wabash Heritage Trail beside the Battlefield monument — wooded creek-side trail, no bikes.",
     hrs("7am-9pm","7am-9pm","7am-9pm","7am-9pm","7am-9pm","7am-9pm","7am-9pm"), "", "",
     "place_id ChIJkYkh89sDE4gRxQ0-DWkF-g0", anchor="approximate"),
 row("Adventures", "Celery Bog Nature Area", "Adventures; Nature Area; Wetland; Birding; Trails; City Park",
     "1620 Lindberg Rd, West Lafayette, IN 47906", 40.4514907, -86.9356462,
     "West Lafayette's wetland nature area with paved and unpaved trails and the Lilly Nature Center; a birding spot.",
     hrs("6am-8:30pm","6am-8:30pm","6am-8:30pm","6am-8:30pm","6am-8:30pm","6am-8:30pm","6am-8:30pm"), "https://www.westlafayette.in.gov/government/parks-and-recreation/places/celery-bog-nature-area", "+1 765-775-5172",
     "place_id ChIJR2FMfST9EogRXOE5hFdLXRI"),
 row("Adventures", "Bicentennial Nature Area", "Adventures; Nature Area; Lake; Fishing; Paddling; Trails",
     "Off US 52 (N of Schuyler Ave), Lafayette, IN 47905", 40.4498615, -86.8729828,
     "Former quarry lake turned nature area on Lafayette's north side — shoreline trail, fishing, a small boat ramp for paddling.",
     hrs("7am-9pm","7am-9pm","7am-9pm","7am-9pm","7am-9pm","7am-9pm","7am-9pm"), "", "+1 765-807-1500",
     "Google gives no street address; anchor approximate; place_id ChIJW8bTpcICE4gRgLo67IjEeTA", anchor="approximate"),
 row("Adventures", "Hoffman Nature Area", "Adventures; Nature Area; Mountain Biking; Hiking; Trails",
     "6935 Old State Rd 25, Lafayette, IN 47905", 40.517304, -86.77467,
     "Hilly nature area east of Prophetstown with a purpose-built mountain-bike trail system; hiking welcome.",
     hrs("7am-8:30pm","7am-8:30pm","7am-8:30pm","7am-8:30pm","7am-8:30pm","7am-8:30pm","7am-8:30pm"), "", "",
     "place_id ChIJyZraGHsGE4gRygNBmvfteXo"),
 row("Adventures", "Prophetstown Fen Nature Preserve", "Adventures; Nature Preserve; Fen; State Dedicated",
     "Inside Prophetstown State Park, West Lafayette, IN 47906", 40.5122598, -86.8188605,
     "State-dedicated nature preserve within Prophetstown State Park protecting hillside seeps and fens above the Wabash floodplain.",
     NONE, "https://www.in.gov/dnr/nature-preserves/", "",
     "within the state park (park gate fee/hours apply); Google address is a parcel number; anchor approximate; place_id ChIJHa_J1tkHE4gRlqSaquv6B-0", anchor="approximate"),
 row("Adventures", "Michaud-Sinninger Woods Nature Preserve", "Adventures; Nature Preserve; Woods; Trails",
     "Cumberland Park, West Lafayette, IN 47906", 40.4654557, -86.9197005,
     "Small wooded preserve tucked into Cumberland Park — a half-mile of wood-chip trail, pond and marsh edge. Park at Cumberland Park.",
     ALL, "", "", "place_id ChIJPadNdTP9EogRG3ZA6J748Gk", anchor="approximate"),
 row("Adventures", "Greenway River Access", "Adventures; River Access; Boat Ramp; Canoe; Kayak; Wabash River",
     "3301 N River Rd, West Lafayette, IN 47906", 40.4639115, -86.8885748,
     "Public ramp and canoe/kayak launch on the Wabash north of West Lafayette; the ramp silts up after floods.",
     NONE, "", "", "place_id ChIJtZtlFgADE4gRpoK_J9C5cT0"),
 row("Adventures", "Davis Ferry Park", "Adventures; County Park; Wabash River; Historic Bridge; Trail; Fishing",
     "N 9th St Rd, Lafayette, IN 47904", 40.4733713, -86.870914,
     "Riverside county park at the old Davis Ferry bridge (now a pedestrian crossing) — river access, fishing, and a link in the Wabash Heritage Trail.",
     NONE, "", "", "place_id ChIJ8VGmDFcCE4gRE0hR1x46TmE"),
 row("Adventures", "Wildcat Creek Park", "Adventures; Park; Creek; Wading; Paddling; Fishing; DNR",
     "5201 Eisenhower Rd, Lafayette, IN 47905", 40.4387798, -86.8038308,
     "Creek-side park on Wildcat Creek east of Lafayette — shallow wading, paddling take-out, fishing and picnicking; now DNR-managed.",
     hrs("8am-9:30pm","8am-8:30pm","8am-8:30pm","8am-8:30pm","8am-9pm","8am-8:30pm","8am-8:30pm"), "", "+1 765-589-8081",
     "place_id ChIJDQdVKLgEE4gRbuNSZr2cNyE"),
 row("Adventures", "Mis-So-Lah Public Access", "Adventures; River Access; Wildcat Creek; Paddling; Tubing; NICHES",
     "Wildcat Creek near CR 600 E, Lafayette, IN 47905", 40.444313, -86.761795,
     "NICHES Land Trust public access on Wildcat Creek upstream of Wildcat Creek Park — put-in for paddling and tubing down to the park.",
     NONE, "https://www.nicheslandtrust.org/", "", "Google gives no street address; anchor approximate; place_id ChIJa-A7dz8FE4gRGjlxjc-SrDQ", anchor="approximate"),
 row("Adventures", "Tapawingo Park", "Adventures; City Park; Riverfront; Wabash River; Events; Ice Rink",
     "100 Tapawingo Dr, West Lafayette, IN 47906", 40.4206831, -86.8988519,
     "West Lafayette's riverfront park at the pedestrian bridge — festival ground, Wabash Heritage Trail access, winter ice rink.",
     hrs("7am-8pm","7am-8pm","7am-8pm","7am-8pm","7am-8pm","7am-8pm","7am-8pm"), "https://www.westlafayette.in.gov/", "+1 765-775-5100",
     "place_id ChIJa3tYYFcdE4gRT6_9hdpmzPs"),
 row("Adventures", "Monon High Bridge Trail", "Adventures; Trail; Rail Trail; Deer Creek; Historic Bridge",
     "1030 N Washington St, Delphi, IN 46923", 40.594766, -86.6561132,
     "Paved rail-trail out of Delphi to the Monon High Bridge over Deer Creek — the first half rebuilt as a walkway with an overlook; the old trestle beyond is closed.",
     NONE, "https://www.cityofdelphi.org/", "+1 765-564-6757", "place_id ChIJE4z_0V2nE4gR7IKJ-2sS06E"),
 row("Adventures", "Altherr Nature Park", "Adventures; Nature Park; Trails; Gardens; City Park",
     "724 S Main St, Monticello, IN 47960", 40.7347684, -86.7618928,
     "Monticello's nature park beside the Bluestem Nature Center — gravel loop trails, gardens, labelled trees.",
     NONE, "https://www.monticelloin.gov/", "+1 574-583-3983", "place_id ChIJA8JsosGvE4gREEH-xu1u8Uk"),
 row("Adventures", "Bluestem Nature Center", "Adventures; Nature Center; Education; Birding; Accessible",
     "802 S Main St, Monticello, IN 47960", 40.7342332, -86.7625148,
     "City-run nature center in Monticello with rotating seasonal exhibits, a bird-watching room and programs for all ages; single-storey, accessible.",
     hrs("Closed","Closed","10am-3:30pm","10am-3:30pm","10am-3:30pm","Closed","Closed"), "https://www.monticelloin.gov/", "+1 574-583-3983",
     "hours as listed on Google today; place_id ChIJySuyFUCvE4gR3oxHuIurGAw"),
 row("Adventures", "Monticello City Park", "Adventures; City Park; Trails; Creek; Ball Fields; Shelters",
     "904 City Park Loop, Monticello, IN 47960", 40.7310232, -86.7601781,
     "Monticello's main city park — creek-side trails, Rotary Cove, ball fields, shelters and a holiday light display.",
     NONE, "https://www.monticelloin.gov/", "+1 574-583-3983", "place_id ChIJ0QJpvsSvE4gROzegnzVBxBs"),
 row("Adventures", "NICHES Land Trust", "Adventures; Land Trust; Nonprofit; Nature Preserves; Volunteer",
     "1782 N 400 E, Lafayette, IN 47905", 40.444519, -86.829123,
     "Regional land trust protecting and opening natural areas across west-central Indiana — preserves with public trails (Wildcat Creek, Wabash and Tippecanoe watersheds).",
     NONE, "https://www.nicheslandtrust.org/", "+1 765-423-1605",
     "Google lists 8am–8pm daily, which reads as preserve hours not office hours — left blank; the preserves deserve their own rows later; place_id ChIJleC4kVodE4gRp1sqx27XKz8"),

 # ---------------- Arts & Entertainment (history)
 row("Arts & Entertainment", "Tippecanoe Battlefield Museum", "Arts & Entertainment; Museum; History; Battle of Tippecanoe; Nature Center",
     "200 Battleground Ave, Battle Ground, IN 47920", 40.5068922, -86.8438448,
     "Museum on the 1811 battlefield — the battle, Prophetstown, the Nations of the Wabash, and the fur-trade era; also the park's nature center. Run by the Tippecanoe County Historical Association. Admission.",
     hrs("10am-5pm","10am-5pm","Closed","10am-5pm","10am-5pm","10am-5pm","10am-5pm"), "https://www.tippecanoehistory.org/", "+1 765-567-2147",
     "hours as listed on Google today — confirm on tippecanoehistory.org; place_id ChIJGSq7vaADE4gR8e4WAEXHarw"),
 row("Arts & Entertainment", "Tippecanoe County Historical Association", "Arts & Entertainment; History; Archives; Genealogy; Research Library",
     "1001 South St, Lafayette, IN 47901", 40.4170027, -86.8855449,
     "TCHA's offices, archives and research library — local history, genealogy, historic markers, and the organisation behind the Battlefield Museum, Fort Ouiatenon and the Feast of the Hunters' Moon.",
     hrs("Closed","Closed","Closed","1pm-5pm","10am-4pm","Closed","Closed"), "https://www.tippecanoehistory.org/", "+1 765-476-8411",
     "research hours as listed on Google today; place_id ChIJQbyY8EYdE4gR4eRPMGUQKO4"),
]

wb = load_workbook(WB)
def norm(s): return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()
added = 0; skipped = []
for r in ROWS:
    ws = wb[r["sheet"]]
    headers = [c.value for c in ws[1]]
    existing = {norm(row[headers.index("Name")]) for row in ws.iter_rows(min_row=2, values_only=True) if row[headers.index("Name")]}
    if norm(r["Name"]) in existing:
        skipped.append(r["Name"]); continue
    ws.append([r.get(h, None) if r.get(h, None) not in ("",) else None for h in headers])
    added += 1
wb.save(WB)
print(f"added {added}; skipped (existing) {skipped}")
