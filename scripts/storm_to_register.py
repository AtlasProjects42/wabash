#!/usr/bin/env python3
"""Convert NCEI Storm Events CSVs into data/climate_events.xlsx (the Regional Climate & Weather Event
Register) in the schema build.py's emit_phenology_history() reads: a README sheet, an "Events" sheet with
the kit's columns, and a Sources sheet. Filters to a county whitelist (the Wabash region). No significance
filter (per Laurie, 2026-10-06): every region event is kept, each traceable to its NCEI EVENT_ID. Run:
    python3 -I scripts/storm_to_register.py <dir-with-storm_data*.csv> --out data/climate_events.xlsx
"""
import csv, glob, os, re, sys
from openpyxl import Workbook

args=sys.argv[1:]
OUT="data/climate_events.xlsx"
if "--out" in args: i=args.index("--out"); OUT=args[i+1]; del args[i:i+2]
SRCDIR=args[0] if args else "."

REGION={'TIPPECANOE','CARROLL','WHITE','CLINTON','WARREN','BENTON','MONTGOMERY','FOUNTAIN','CASS','HOWARD','TIPTON','BOONE','JASPER','NEWTON'}

def clean_cz(s):
    return re.sub(r'\s*\((ZONE|CO\.?)\)|\s*CO\.?$','',(s or '').upper()).strip()
def iso(d):
    d=(d or '').strip()
    m=re.match(r'(\d{1,2})/(\d{1,2})/(\d{4})',d)
    if m: return f"{int(m.group(3)):04d}-{int(m.group(1)):02d}-{int(m.group(2)):02d}"
    m=re.match(r'(\d{4})-(\d{2})-(\d{2})',d)
    if m: return d[:10]
    return ""
def num(v):
    try: return float(v)
    except (TypeError,ValueError): return 0.0
def money(v):
    n=num(v)
    if n<=0: return ""
    if n>=1e6: return f"${n/1e6:.1f}M"
    if n>=1e3: return f"${n/1e3:.0f}K"
    return f"${n:.0f}"

def key_measure(r):
    et=r.get('EVENT_TYPE','') or ''; parts=[]
    if 'Tornado' in et:
        f=(r.get('TOR_F_SCALE','') or '').strip()
        if f: parts.append(f)
        L=num(r.get('TOR_LENGTH')); W=num(r.get('TOR_WIDTH'))
        if L>0: parts.append(f"{L:g} mi path")
        if W>0: parts.append(f"{W:g} yd wide")
    elif et=='Hail':
        m=num(r.get('MAGNITUDE'));
        if m>0: parts.append(f"{m:g} in hail")
    elif 'Wind' in et:
        m=num(r.get('MAGNITUDE')); mt=(r.get('MAGNITUDE_TYPE','') or '').strip()
        if m>0: parts.append(f"{m:g} kt" + (f" ({mt})" if mt else ""))
    elif et in ('Flood','Flash Flood'):
        fc=(r.get('FLOOD_CAUSE','') or '').strip()
        if fc: parts.append(fc)
    else:
        m=num(r.get('MAGNITUDE'))
        if m>0: parts.append(f"{m:g}")
    return " · ".join(parts)

def impact(r):
    d=int(num(r.get('DEATHS_DIRECT'))+num(r.get('DEATHS_INDIRECT')))
    i=int(num(r.get('INJURIES_DIRECT'))+num(r.get('INJURIES_INDIRECT')))
    dmg=money(str(num(r.get('DAMAGE_PROPERTY_NUM'))+num(r.get('DAMAGE_CROPS_NUM'))))
    head=[]
    if d: head.append(f"{d} killed")
    if i: head.append(f"{i} injured")
    if dmg: head.append(f"{dmg} damage")
    narr=(r.get('EVENT_NARRATIVE','') or '').strip() or (r.get('EPISODE_NARRATIVE','') or '').strip()
    pieces=[]
    if head: pieces.append("; ".join(head)+".")
    if narr: pieces.append(narr)
    return " ".join(pieces)

def title(r):
    et=r.get('EVENT_TYPE','') or 'Weather event'
    place=(r.get('BEGIN_LOCATION','') or '').strip()
    cz=clean_cz(r.get('CZ_NAME_STR',''))
    where = (place.title()+f" ({cz.title()} Co.)") if place else (cz.title()+" Co.")
    return f"{et} — {where}"

rows=[]
for f in sorted(glob.glob(os.path.join(SRCDIR,'*storm_data*.csv'))):
    with open(f,newline='',encoding='utf-8',errors='replace') as fh:
        for r in csv.DictReader(fh):
            if clean_cz(r.get('CZ_NAME_STR','')) in REGION:
                rows.append(r)
# de-dup by EVENT_ID (a file set can overlap)
seen=set(); uniq=[]
for r in rows:
    eid=r.get('EVENT_ID','')
    if eid in seen: continue
    seen.add(eid); uniq.append(r)
rows=uniq
rows.sort(key=lambda r:(iso(r.get('BEGIN_DATE','')), r.get('BEGIN_TIME','')))

HEADERS=['ID','Start date','End date','Year','Century','Category','Event name','Area affected','Key measurement',
 'Station / gauge','Impact summary','Primary source','Source URL','Confidence','Verified this pass','Notes / to-do',
 'Timeline anchor (ISO)','Anchor year','Anchor month','Anchor day','Date precision','Anchor basis','Show on timeline']

wb=Workbook()
rd=wb.active; rd.title='README'
for line in [
 ['Regional Climate & Weather Event Register — Wabash country around Prophetstown, Indiana'],
 ['Built from NOAA NCEI Storm Events Database (storm_data CSVs), '+__import__('datetime').date.today().isoformat()+'.'],
 ['Scope: the Wabash region — '+", ".join(sorted(c.title() for c in REGION))+' counties (Laurie, 2026-10-06).'],
 ['Filter: none — every NCEI Storm Events record in these counties is included. Each row keeps its NCEI EVENT_ID (ID = SE-<id>) and a link to its event page.'],
 ['Read by build.py emit_phenology_history() → the "On record" timeline on Signs + Signals.'],
 ['Add non-NCEI historical events (pre-1950 floods, newspaper-sourced storms) as new rows in the Events sheet, same columns.'],
]: rd.append(line)

ws=wb.create_sheet('Events'); ws.append(HEADERS)
for r in rows:
    s=iso(r.get('BEGIN_DATE','')); e=iso(r.get('END_DATE','')) or s
    if not s: continue
    y=int(s[:4]); cat=r.get('EVENT_TYPE','') or 'Weather event'
    ws.append([
      'SE-'+(r.get('EVENT_ID','') or ''), s, e, y, (y-1)//100+1, cat, title(r),
      (clean_cz(r.get('CZ_NAME_STR','')).title()+' Co., IN'), key_measure(r), '', impact(r),
      'NOAA NCEI Storm Events Database',
      'https://www.ncdc.noaa.gov/stormevents/eventdetails.jsp?id='+(r.get('EVENT_ID','') or ''),
      'A', '', '', s, y, int(s[5:7]), int(s[8:10]), 'day',
      ('Single-day event' if s==e else 'Multi-day event'), 'Yes'
    ])

sc=wb.create_sheet('Sources'); sc.append(['Source','URL','Best for'])
sc.append(['NOAA NCEI Storm Events Database','https://www.ncdc.noaa.gov/stormevents/','Official NWS record of tornadoes, floods, severe storms, winter storms and more, 1950–present, by county.'])

os.makedirs(os.path.dirname(OUT) or '.', exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}: {len(rows)} region events ({rows[0] and iso(rows[0].get('BEGIN_DATE','')) } → {iso(rows[-1].get('BEGIN_DATE',''))})")
import collections
print("by type:", dict(collections.Counter(r.get('EVENT_TYPE','') for r in rows).most_common(12)))
